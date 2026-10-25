"""Build the display-only Excel template rtm_display.xlsx (rtmtools, schema rtm-export-1.0).

The template is copied into every run folder by the Stata export helper (_rtm_export), which
fills the data sheets (manifest and one sheet per export table) with export excel. Every cell
on the display sheets is a reference or a lookup into those data sheets, or a chart of them.
Nothing is computed here: no probability, weight, factor, CV or fit is derived from another
number. The allowed functions are listed on the Notes sheet and enforced by tests/check_display.py.

Run from the package folder:  python3 excel/build_display.py   -> writes rtm_display.xlsx at the root
"""
import os
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference, ScatterChart, Series
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, "rtm_display.xlsx")

# the export tables and their columns (schema rtm-export-1.0)
TABLES = {
    "frame_units": ["unit_id", "stratum", "unit", "size", "x", "y", "flag"],
    "frame_summary": ["stratum", "n_units", "size_sum", "size_min", "size_p50", "size_max", "n_flagged"],
    "boundaries": ["polygon_id", "part", "vertex", "x", "y"],
    "listing_areas": ["area_id", "stratum", "expected", "listed", "eligible", "ratio", "segments", "segment_selected", "flag"],
    "listing_checks": ["check", "n", "status", "detail"],
    "size": ["domain", "p", "margin", "deff", "nonresponse", "attrition", "rounds", "n_effective", "n_final"],
    "allocation": ["stratum", "n_units", "size_sum", "share", "n_planned", "n_min", "n_max", "n_final", "sampling_fraction"],
    "selection": ["unit_id", "stratum", "size", "order", "key", "cum_size", "selection_number", "selected", "reserve", "rank", "prob", "weight", "x", "y"],
    "draw_summary": ["stratum", "n_units", "n_selected", "n_reserve", "prob_min", "prob_max", "interval", "random_start"],
    "weight_stages": ["stage", "variable", "n", "sum", "min", "p1", "p5", "p25", "p50", "p75", "p95", "p99", "max", "mean", "cv", "kish_deff", "n_trimmed", "n_zero"],
    "weight_groups": ["stage", "variable", "group_var", "group", "n", "sum", "mean", "cv"],
    "response_classes": ["class_var", "class", "n_sample", "n_respondents", "response_rate", "factor"],
    "calibration_fit": ["constraint", "target", "before", "after", "diff", "rel_diff", "converged", "iterations"],
    "weight_units": ["unit_id", "stratum", "psu", "flag"],
    "flags": ["table", "unit_id", "stage", "value", "reason"],
    "replacements": ["unit_id", "candidate_id", "rank", "distance", "stratum", "adjacency"],
}
ALLOWED_FUNCTIONS = ["IF", "IFERROR", "INDEX", "MATCH", "COUNTIF", "COUNTA", "NA"]
NOT_IN = "not in this export"

HEAD = Font(bold=True, size=14)
SUB = Font(bold=True, size=11)
MUTED = Font(italic=True, color="666666")
FILL = PatternFill("solid", fgColor="EEF4FC")


def manifest_lookup(key, default=NOT_IN):
    # INDEX/MATCH inside IFERROR: a lookup that every Excel version evaluates (XLOOKUP needs the
    # _xlfn prefix when written outside Excel and is absent from Excel 2016/2019)
    return f'=IFERROR(INDEX(manifest!$B:$B,MATCH("{key}",manifest!$A:$A,0)),"{default}")'


def ref_block(ws, top, left, sheet, cols, nrows, fmt=None):
    """A block that mirrors the first nrows rows of a data sheet: =IF(src="","",src)."""
    for j, col in enumerate(cols):
        c = ws.cell(row=top, column=left + j, value=col)
        c.font = SUB
        c.fill = FILL
    for i in range(nrows):
        for j in range(len(cols)):
            src = f"{sheet}!{get_column_letter(j + 1)}{i + 2}"
            c = ws.cell(row=top + 1 + i, column=left + j, value=f'=IF({src}="","",{src})')
            if fmt and j in fmt:
                c.number_format = fmt[j]
    return top + 1 + nrows


def run_block(ws, r):
    ws.cell(row=r, column=1, value="Run").font = SUB
    for key, label in [("run_id", "Run id"), ("step", "Step"), ("started", "Started"), ("status", "Status"),
                       ("source", "Source (ado or do-file)"), ("input_file", "Input file"), ("input_rows", "Input rows"), ("command", "Command")]:
        r += 1
        ws.cell(row=r, column=1, value=label)
        ws.cell(row=r, column=2, value=manifest_lookup(key)).alignment = Alignment(horizontal="left")
    return r


def keys_block(ws, r, title, keys):
    r += 2
    ws.cell(row=r, column=1, value=title).font = SUB
    for key, label in keys:
        r += 1
        ws.cell(row=r, column=1, value=label)
        c = ws.cell(row=r, column=2, value=manifest_lookup(key))
        c.number_format = "#,##0.000"
        c.alignment = Alignment(horizontal="right")
    return r


def bar(ws, title, col_idx, top, nrows, anchor, cat_col=1, ytitle=None):
    ch = BarChart()
    ch.type = "col"
    ch.title = title
    if ytitle:
        ch.y_axis.title = ytitle
    ch.add_data(Reference(ws, min_col=col_idx, min_row=top, max_row=top + nrows), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=cat_col, min_row=top + 1, max_row=top + nrows))
    ch.height, ch.width = 7, 16
    ch.legend = None
    ws.add_chart(ch, anchor)


def pair_bar(ws, title, col_a, col_b, top, nrows, anchor, cat_col=1):
    ch = BarChart()
    ch.type = "col"
    ch.title = title
    ch.add_data(Reference(ws, min_col=col_a, min_row=top, max_row=top + nrows), titles_from_data=True)
    ch.add_data(Reference(ws, min_col=col_b, min_row=top, max_row=top + nrows), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=cat_col, min_row=top + 1, max_row=top + nrows))
    ch.height, ch.width = 7, 16
    ws.add_chart(ch, anchor)


def scatter_helpers(ws, top, left, sheet, xcol, ycol, flagcol, flagval, label, nrows):
    """Two helper columns: the x and y of the rows whose flag equals flagval, NA() otherwise (NA keeps
    a point off a scatter chart; it computes nothing)."""
    ws.cell(row=top, column=left, value=f"{label}: x").font = SUB
    ws.cell(row=top, column=left + 1, value=f"{label}: y").font = SUB
    for i in range(nrows):
        rr = i + 2
        cond = f"{sheet}!{flagcol}{rr}={flagval}" if flagval is not None else f"{sheet}!{xcol}{rr}<>\"\""
        ws.cell(row=top + 1 + i, column=left, value=f'=IF({cond},{sheet}!{xcol}{rr},NA())')
        ws.cell(row=top + 1 + i, column=left + 1, value=f'=IF({cond},{sheet}!{ycol}{rr},NA())')


def scatter(ws, title, series, anchor):
    ch = ScatterChart()
    ch.title = title
    ch.style = 13
    for (xref, yref, name) in series:
        s = Series(yref, xref, title=name)
        s.marker.symbol = "circle"
        s.marker.size = 5
        s.graphicalProperties.line.noFill = True
        ch.series.append(s)
    ch.height, ch.width = 10, 14
    ws.add_chart(ch, anchor)


def build_sampling(wb):
    w = wb.create_sheet("Sampling")
    w.column_dimensions["A"].width = 30
    for col in "BCDEFGHIJKLMNOPQRS":
        w.column_dimensions[col].width = 13
    w["A1"] = "Sampling: what Stata exported for this run"
    w["A1"].font = HEAD
    w["A2"] = "Sizes, allocation and the draw as computed in Stata; nothing is recalculated here (see Notes)."
    w["A2"].font = MUTED
    r = run_block(w, 4)
    r = keys_block(w, r, "Key numbers (r() results of the command)", [
        ("summary_n_domain", "Households per domain (size)"), ("summary_n_total", "Households in all domains"), ("summary_n_srs", "Interviews, simple random sample"),
        ("summary_f_attrition", "Attrition reserve factor"), ("summary_n", "Total to allocate"), ("summary_n_rounded", "Allocated after rounding"),
        ("summary_N_frame", "Frame units"), ("summary_n_selected", "Selected units"), ("summary_n_reserve", "Reserve units"), ("summary_n_certainty", "Certainty units"),
        ("summary_H", "Strata"), ("summary_seed", "Seed"), ("summary_bands", "Serpentine bands"), ("summary_n_numbers", "Numbers generated (RDD)")])
    r += 2
    w.cell(row=r, column=1, value="Sample size per domain (size.csv)").font = SUB
    r += 1
    r = ref_block(w, r, 1, "size", TABLES["size"], 6, {1: "0.000", 2: "0.0000", 3: "0.00", 4: "0.00", 5: "0.000", 6: "0", 7: "#,##0.0", 8: "#,##0"})
    r += 1
    w.cell(row=r, column=1, value="Allocation across strata (allocation.csv): planned against final").font = SUB
    r += 1
    top = r
    r = ref_block(w, r, 1, "allocation", TABLES["allocation"], 40, {1: "#,##0", 2: "#,##0", 3: "0.000", 4: "0.00", 5: "0", 6: "0", 7: "#,##0", 8: "0.0000"})
    pair_bar(w, "Planned and final allocation by stratum (allocation.csv)", 5, 8, top, 40, f"K{top}")
    r += 1
    w.cell(row=r, column=1, value="Draw by stratum (draw_summary.csv)").font = SUB
    r += 1
    top = r
    r = ref_block(w, r, 1, "draw_summary", TABLES["draw_summary"], 40, {1: "#,##0", 2: "#,##0", 3: "#,##0", 4: "0.0000", 5: "0.0000", 6: "#,##0.0", 7: "#,##0.0"})
    pair_bar(w, "Inclusion probability, min and max by stratum (draw_summary.csv)", 5, 6, top, 40, f"K{top}")
    r += 1
    w.cell(row=r, column=1, value="Selected and reserve units (selection.csv, first 60 rows; the file holds all)").font = SUB
    r += 1
    r = ref_block(w, r, 1, "selection", TABLES["selection"], 60, {2: "#,##0", 3: "0", 4: "0.000000", 5: "#,##0", 6: "#,##0.0", 7: "0", 8: "0", 9: "0", 10: "0.0000", 11: "#,##0.00", 12: "#,##0", 13: "#,##0"})
    # the map: selected units against the frame, from the exported coordinates (helper columns U:X)
    scatter_helpers(w, 4, 21, "selection", "M", "N", "H", 1, "selected", 2000)
    scatter_helpers(w, 4, 23, "selection", "M", "N", "H", 0, "not selected", 2000)
    scatter(w, "Selected units (dark) and the rest of the frame (selection.csv x, y)", [
        (Reference(w, min_col=23, min_row=5, max_row=2004), Reference(w, min_col=24, min_row=5, max_row=2004), "frame"),
        (Reference(w, min_col=21, min_row=5, max_row=2004), Reference(w, min_col=22, min_row=5, max_row=2004), "selected")], "K70")
    w.freeze_panes = "A4"


def build_frame(wb):
    w = wb.create_sheet("Frame")
    w.column_dimensions["A"].width = 30
    for col in "BCDEFGHIJKLMNOPQRS":
        w.column_dimensions[col].width = 13
    w["A1"] = "Frame: what Stata exported for this run"
    w["A1"].font = HEAD
    w["A2"] = "List-frame checks or the grid frame as computed in Stata; nothing is recalculated here (see Notes)."
    w["A2"].font = MUTED
    r = run_block(w, 4)
    r = keys_block(w, r, "Key numbers (r() results of the command)", [
        ("summary_N", "Units"), ("summary_N_ok", "Usable units"), ("summary_n_dup", "Duplicate identifiers"), ("summary_n_badsize", "Bad sizes"),
        ("summary_n_small", "Small units or cells"), ("summary_n_large", "Large units"), ("summary_n_excluded", "Excluded"), ("summary_H", "Strata"),
        ("summary_N_cells", "Grid cells"), ("summary_n_polygons", "Boundary polygons"), ("summary_n_straddle", "Cells straddling a boundary"),
        ("summary_pop_inside", "Population inside the boundaries"), ("summary_pop_outside", "Population outside"), ("summary_pixels_outside", "Pixels outside"),
        ("summary_cellsize", "Cell size (m)"), ("summary_pixelsize", "Pixel size (m)"), ("summary_n_with_candidates", "Cells with replacement candidates"), ("summary_n_parts", "Parts after a split")])
    r += 2
    w.cell(row=r, column=1, value="Summary by stratum (frame_summary.csv)").font = SUB
    r += 1
    top = r
    r = ref_block(w, r, 1, "frame_summary", TABLES["frame_summary"], 40, {1: "#,##0", 2: "#,##0", 3: "#,##0", 4: "#,##0", 5: "#,##0", 6: "#,##0"})
    bar(w, "Size (population or households) by stratum (frame_summary.csv)", 3, top, 40, f"J{top}")
    bar(w, "Flagged units by stratum (frame_summary.csv)", 7, top, 40, f"J{top + 15}")
    r += 1
    w.cell(row=r, column=1, value="Replacement candidates (replacements.csv)").font = SUB
    r += 1
    r = ref_block(w, r, 1, "replacements", TABLES["replacements"], 30, {2: "0", 3: "#,##0"})
    r += 1
    w.cell(row=r, column=1, value="Units or cells (frame_units.csv, first 60 rows; the file holds all)").font = SUB
    r += 1
    r = ref_block(w, r, 1, "frame_units", TABLES["frame_units"], 60, {3: "#,##0", 4: "#,##0", 5: "#,##0"})
    scatter_helpers(w, 4, 21, "frame_units", "E", "F", None, None, "units", 2000)
    scatter(w, "Units or cells of the frame (frame_units.csv x, y)", [
        (Reference(w, min_col=21, min_row=5, max_row=2004), Reference(w, min_col=22, min_row=5, max_row=2004), "units")], "J62")
    w.freeze_panes = "A4"


def build_listing(wb):
    w = wb.create_sheet("Listing")
    w.column_dimensions["A"].width = 44
    for col in "BCDEFGHIJKLMNOPQRS":
        w.column_dimensions[col].width = 13
    w["A1"] = "Listing: what Stata exported for this run"
    w["A1"].font = HEAD
    w["A2"] = "The listing forms and the completed listing's checks as computed in Stata; nothing is recalculated here (see Notes)."
    w["A2"].font = MUTED
    r = run_block(w, 4)
    r = keys_block(w, r, "Key numbers (r() results of the command)", [
        ("summary_N_areas", "Selected areas"), ("summary_n_segmented", "Areas segmented"), ("summary_N_rows", "Rows in the listing"),
        ("summary_n_returned", "Areas returned"), ("summary_n_not_returned", "Areas not returned"), ("summary_n_not_selected", "Households in areas not selected"),
        ("summary_n_duplicates", "Duplicate rows"), ("summary_n_under", "Under-listed areas"), ("summary_n_over", "Over-listed areas"), ("summary_N_frame", "Eligible households in the frame")])
    r += 2
    w.cell(row=r, column=1, value="Checks (listing_checks.csv)").font = SUB
    r += 1
    r = ref_block(w, r, 1, "listing_checks", TABLES["listing_checks"], 12, {1: "#,##0"})
    r += 1
    w.cell(row=r, column=1, value="Areas (listing_areas.csv, first 100 rows)").font = SUB
    r += 1
    top = r
    r = ref_block(w, r, 1, "listing_areas", TABLES["listing_areas"], 100, {2: "#,##0", 3: "#,##0", 4: "#,##0", 5: "0.000", 6: "0", 7: "0"})
    scatter_helpers(w, 4, 21, "listing_areas", "C", "D", None, None, "areas", 500)
    scatter(w, "Listed against expected households by area (listing_areas.csv)", [
        (Reference(w, min_col=21, min_row=5, max_row=504), Reference(w, min_col=22, min_row=5, max_row=504), "areas")], f"K{top}")
    w.freeze_panes = "A4"


def build():
    wb = Workbook()
    w = wb.active
    w.title = "Weights"
    w.column_dimensions["A"].width = 34
    for col in "BCDEFGHIJKLMNOPQRS":
        w.column_dimensions[col].width = 13
    w["A1"] = "Weights: what Stata exported for this run"
    w["A1"].font = HEAD
    w["A2"] = "Every number on this sheet is read from the manifest and the table sheets of the export; nothing is recalculated here (see Notes)."
    w["A2"].font = MUTED

    # run block
    r = 4
    w.cell(row=r, column=1, value="Run").font = SUB
    for key, label in [("run_id", "Run id"), ("workflow", "Workflow"), ("step", "Step"), ("started", "Started"),
                       ("finished", "Finished"), ("status", "Status"), ("source", "Source (ado or do-file)"),
                       ("software", "Software"), ("input_file", "Input file"), ("input_rows", "Input rows"), ("command", "Command")]:
        r += 1
        w.cell(row=r, column=1, value=label)
        c = w.cell(row=r, column=2, value=manifest_lookup(key))
        c.alignment = Alignment(horizontal="left")
    # settings block
    r += 2
    w.cell(row=r, column=1, value="Settings (as recorded by Stata, with the source of the default)").font = SUB
    for key, label in [("setting_method", "Method"), ("source_method", "Method source"), ("setting_covariates", "Covariates"),
                       ("setting_classes", "Classes"), ("setting_cells", "Cells"), ("setting_moments", "Moments"),
                       ("setting_ngroups", "Propensity groups"), ("setting_rule", "Trimming rule"), ("setting_pct", "Trimming percentile"),
                       ("setting_engine", "Calibration engine"), ("setting_calibration_variables", "Calibration variables"),
                       ("setting_bounds", "Bounds"), ("setting_weight", "Input weight"), ("setting_generate", "New weight")]:
        r += 1
        w.cell(row=r, column=1, value=label)
        w.cell(row=r, column=2, value=manifest_lookup(key)).alignment = Alignment(horizontal="left")
    # key numbers block
    r += 2
    w.cell(row=r, column=1, value="Key numbers (r() results of the command)").font = SUB
    keys = [("summary_N_sample", "Eligible sample"), ("summary_N_resp", "Respondents"), ("summary_rr_unweighted", "Response rate, unweighted"),
            ("summary_rr_weighted", "Response rate, weighted"), ("summary_sum_base", "Sum of base weights"), ("summary_sum_adj", "Sum of adjusted weights"),
            ("summary_factor_min", "Adjustment factor, min"), ("summary_factor_p50", "Adjustment factor, median"), ("summary_factor_max", "Adjustment factor, max"),
            ("summary_max_rel_diff", "Largest relative difference to the targets"), ("summary_ebalance_maxreldiff", "ebalance cross-check, max relative difference"),
            ("summary_n_trimmed_high", "Weights lowered by trimming"), ("summary_n_trimmed_low", "Weights raised by trimming"),
            ("summary_kish_before", "Kish 1+CV2 before"), ("summary_kish_after", "Kish 1+CV2 after"), ("summary_kish", "Kish 1+CV2"),
            ("summary_cv", "CV of the weights"), ("summary_N_pop", "Population total"), ("summary_population_uncovered", "Population in cells with no respondent"),
            ("summary_N_hh", "Households"), ("summary_sum_pop", "Sum of population weights"), ("summary_sum_hh", "Sum of household weights")]
    for key, label in keys:
        r += 1
        w.cell(row=r, column=1, value=label)
        c = w.cell(row=r, column=2, value=manifest_lookup(key))
        c.number_format = "#,##0.000"
        c.alignment = Alignment(horizontal="right")
    # weight stages block
    r += 2
    w.cell(row=r, column=1, value="Weight stages (weight_stages.csv)").font = SUB
    r += 1
    top_ws = r
    fmt = {i: "#,##0.000" for i in range(2, 18)}
    fmt[2] = "#,##0"
    fmt[16] = "#,##0"
    fmt[17] = "#,##0"
    r = ref_block(w, r, 1, "weight_stages", TABLES["weight_stages"], 12, fmt)
    # charts on the stages: CV and Kish by stage
    ch = BarChart()
    ch.type = "col"
    ch.title = "CV of the weights by stage (weight_stages.csv)"
    ch.y_axis.title = "CV"
    ch.add_data(Reference(wb["Weights"], min_col=15, min_row=top_ws, max_row=top_ws + 12), titles_from_data=True)
    ch.set_categories(Reference(wb["Weights"], min_col=1, min_row=top_ws + 1, max_row=top_ws + 12))
    ch.height, ch.width = 7, 16
    ch.legend = None
    w.add_chart(ch, f"A{r + 1}")
    ch2 = BarChart()
    ch2.type = "col"
    ch2.title = "Kish's 1+CV2 by stage (weight_stages.csv)"
    ch2.add_data(Reference(wb["Weights"], min_col=16, min_row=top_ws, max_row=top_ws + 12), titles_from_data=True)
    ch2.set_categories(Reference(wb["Weights"], min_col=1, min_row=top_ws + 1, max_row=top_ws + 12))
    ch2.height, ch2.width = 7, 16
    ch2.legend = None
    w.add_chart(ch2, f"J{r + 1}")
    r += 16
    # response classes block
    w.cell(row=r, column=1, value="Response by class or group (response_classes.csv)").font = SUB
    r += 1
    top_rc = r
    r = ref_block(w, r, 1, "response_classes", TABLES["response_classes"], 30, {2: "#,##0", 3: "#,##0", 4: "0.000", 5: "0.000"})
    ch3 = BarChart()
    ch3.type = "col"
    ch3.title = "Response rate by class (response_classes.csv)"
    ch3.add_data(Reference(wb["Weights"], min_col=5, min_row=top_rc, max_row=top_rc + 30), titles_from_data=True)
    ch3.set_categories(Reference(wb["Weights"], min_col=2, min_row=top_rc + 1, max_row=top_rc + 30))
    ch3.height, ch3.width = 7, 16
    ch3.legend = None
    w.add_chart(ch3, f"H{top_rc}")
    ch4 = BarChart()
    ch4.type = "col"
    ch4.title = "Adjustment factor by class (response_classes.csv)"
    ch4.add_data(Reference(wb["Weights"], min_col=6, min_row=top_rc, max_row=top_rc + 30), titles_from_data=True)
    ch4.set_categories(Reference(wb["Weights"], min_col=2, min_row=top_rc + 1, max_row=top_rc + 30))
    ch4.height, ch4.width = 7, 16
    ch4.legend = None
    w.add_chart(ch4, f"H{top_rc + 15}")
    r += 1
    # calibration fit block
    w.cell(row=r, column=1, value="Calibration or balancing fit (calibration_fit.csv)").font = SUB
    r += 1
    top_cf = r
    r = ref_block(w, r, 1, "calibration_fit", TABLES["calibration_fit"], 30, {1: "#,##0.000", 2: "#,##0.000", 3: "#,##0.000", 4: "#,##0.000", 5: "0.000000"})
    ch5 = BarChart()
    ch5.type = "col"
    ch5.title = "Target and achieved totals (calibration_fit.csv)"
    ch5.add_data(Reference(wb["Weights"], min_col=2, max_col=4, min_row=top_cf, max_row=top_cf + 30), titles_from_data=True)
    ch5.set_categories(Reference(wb["Weights"], min_col=1, min_row=top_cf + 1, max_row=top_cf + 30))
    ch5.height, ch5.width = 7, 16
    w.add_chart(ch5, f"J{top_cf}")
    r += 1
    # groups block
    w.cell(row=r, column=1, value="Weights by group (weight_groups.csv)").font = SUB
    r += 1
    r = ref_block(w, r, 1, "weight_groups", TABLES["weight_groups"], 40, {4: "#,##0", 5: "#,##0.000", 6: "#,##0.000", 7: "0.000"})
    r += 1
    # flags block
    w.cell(row=r, column=1, value="Flagged units (flags.csv)").font = SUB
    r += 1
    w.cell(row=r, column=1, value="Flagged rows in the export:")
    w.cell(row=r, column=2, value='=COUNTA(flags!$B$2:$B$5000)')
    r += 1
    r = ref_block(w, r, 1, "flags", TABLES["flags"], 40, {3: "#,##0.000"})
    w.freeze_panes = "A4"

    build_sampling(wb)
    build_frame(wb)
    build_listing(wb)

    # ---- Notes sheet ----
    n = wb.create_sheet("Notes")
    n.column_dimensions["A"].width = 120
    lines = [
        ("rtmtools display workbook (schema rtm-export-1.0)", HEAD),
        ("This workbook displays what the Stata commands exported for one run (sheets Weights, Sampling, Frame, Listing; use the one of the run's", None),
        ("workflow). Stata copies this template into the run folder and fills the sheet", None),
        ("'manifest' and one sheet per export table (weight_stages, response_classes, calibration_fit, weight_groups, flags, ...) from the same", None),
        ("CSV files that the dashboard reads. Nothing in this workbook recalculates a statistic: every value on the Weights sheet is a reference", None),
        ("or a lookup into those sheets, and every chart plots a column of them.", None),
        ("", None),
        ("Formulas used, the complete list: " + ", ".join(ALLOWED_FUNCTIONS) + ". IF only blanks out empty rows; IFERROR, INDEX and MATCH", SUB),
        ("fetch a value from the manifest; COUNTIF and COUNTA count exported rows; NA() keeps a point off a scatter chart (the map helper", None),
        ("columns U to X of the Sampling, Frame and Listing sheets). No division, multiplication, power, RAND, PERCENTILE, STDEV, Solver or", None),
        ("Analysis ToolPak anywhere. tests/check_display.py reads every formula and fails on any other function.", None),
        ("", None),
        ("Rounding: values are stored exactly as exported; number formats show three decimals (counts as integers).", None),
        ("'" + NOT_IN + "' means the step did not produce that value or table.", None),
        ("", None),
        ("Source of every table and chart: the CSV file named in its heading, in the run folder next to this workbook.", None),
        ("To look at another run without Stata: Data > Get Data > From Text/CSV (Power Query) can load that run's CSV files into the", None),
        ("sheets of the same name, or paste the CSV contents into them; the Weights sheet then shows them. Nothing leaves this computer.", None),
    ]
    for i, (txt, font) in enumerate(lines, start=1):
        c = n.cell(row=i, column=1, value=txt)
        if font:
            c.font = font

    # ---- manifest sheet ----
    m = wb.create_sheet("manifest")
    m["A1"] = "key"
    m["B1"] = "value"
    m.column_dimensions["A"].width = 30
    m.column_dimensions["B"].width = 90

    # ---- table sheets ----
    for name, cols in TABLES.items():
        s = wb.create_sheet(name)
        for j, col in enumerate(cols, start=1):
            s.cell(row=1, column=j, value=col).font = SUB
            s.column_dimensions[get_column_letter(j)].width = 14
    wb.save(OUT)
    return OUT


if __name__ == "__main__":
    print("written:", build())
