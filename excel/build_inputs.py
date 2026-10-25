"""Build the input workbook template rtm_inputs.xlsx (rtmtools, input schema rtm-input-1.0).

A program fills this one workbook (or a folder of CSV files with the same sheet names) and runs
`rtminput using rtm_inputs.xlsx, check | sampling | weights round(#)`. The Design sheet holds the
method choices as key/value rows; Strata, Frame, Totals, Cells and Prefixes hold the aggregate
inputs; Files points to the household-level data files, which stay outside the workbook.

Example rows come from the package's synthetic data, so the template runs as shipped once the
sample data are in the same folder (net get rtmtools).

Run from the package folder:  python3 excel/build_inputs.py
  -> writes rtm_inputs.xlsx at the root (paths such as rtm_baseline.dta, the net get layout)
  -> writes portable_do_files/rtm_inputs.xlsx (paths such as template/rtm_baseline.dta)
"""
import os
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DATA = os.path.join(ROOT, "data")

HEAD = Font(bold=True, size=14)
SUB = Font(bold=True, size=11)
MUTED = Font(italic=True, color="666666")
FILL = PatternFill("solid", fgColor="EEF4FC")
REQ = PatternFill("solid", fgColor="FFF4D6")

# ---- the Design keys: (key, example value, note with allowed values and the source of the default) ----
DESIGN = [
    ("# project", "", ""),
    ("project", "RTM synthetic example", "free text: the program's name, written into the exports"),
    ("seed", "20261024", "integer; the same seed repeats every draw; blank: Stata draws one and records it"),
    ("output_dir", "rtm_outputs", "folder for the output files, relative to the folder you run Stata from"),
    ("# sample size (rtmsample size)", "", ""),
    ("target", "level", "level | change: what the sample size is set for"),
    ("p", "0.2", "expected proportion of the key indicator (UN 2008 para 59)"),
    ("relmargin", "0.10", "relative margin of error at 95 percent (UN 2008 para 59: 10 percent); or set margin below"),
    ("margin", "", "absolute margin instead of relmargin"),
    ("change", "", "with target change: the change to detect, for example -0.05"),
    ("power", "0.8", "with target change: power of the test"),
    ("deff", "2", "design effect (UN 2008 para 60: about 2 for clustered samples, about 1 for a list sample)"),
    ("nonresponse_mult", "1.1", "nonresponse multiplier (UN 2008 para 60)"),
    ("attrition", "0.015", "attrition per round of the panel (MPPS 2016 p. 45); 0 = none"),
    ("rounds", "24", "rounds of the panel"),
    ("firstloss", "0.25", "loss between the baseline and the first phone round; 0 = none"),
    ("domains", "4", "domains that need equal reliability"),
    ("# frame check and allocation (rtmframe list, rtmsample allocate)", "", ""),
    ("minsize", "80", "areas with fewer households are flagged (GM 1996 pp. 78-79: appended to a neighbour); 0 = none"),
    ("maxsize", "250", "areas with more are flagged for segmentation (UN 2008 paras 138-140); 0 = none"),
    ("n_areas", "100", "areas to select in total (required for sampling)"),
    ("allocation", "proportional", "proportional | equal | neyman | power (UN 2008 paras 68-70; Cochran 1977 5.5; Bankier 1988)"),
    ("power_q", "0.5", "exponent of power allocation"),
    ("min_areas", "10", "minimum areas per stratum; 0 = none"),
    ("max_areas", "16", "maximum areas per stratum; 0 = none"),
    ("# first stage (rtmsample draw or rtmsample grid)", "", ""),
    ("stage1", "ppssys", "ppssys | pps | systematic | srs | grid (UN 2008 paras 77-79 and 118-119; grid = serpentine order, needs x y)"),
    ("sort_vars", "stratum_id area_id", "Frame columns that give the geographic order within strata (implicit stratification)"),
    ("reserve_areas", "2", "reserve areas per stratum: a number, all, or 0 (UN 2008 paras 172, 178-179)"),
    ("# second stage (rtmsample draw within the selected areas)", "", ""),
    ("households_per_area", "12", "households per selected area, from the household_list or listing file of the Files sheet; blank = no second stage"),
    ("# phone subsample of a baseline (rtmsample draw, reserve all)", "", ""),
    ("phone_per_psu", "6", "households per PSU called first, every other household an ordered replacement; blank = none (needs a baseline row in Files)"),
    ("# random digit dialing (rtmsample rdd)", "", ""),
    ("rdd", "no", "yes | no: draw numbers from the Prefixes sheet; the numbers are never exported"),
    ("# weights of a round (rtmweight base, nonresponse, trim, calibrate, report)", "", ""),
    ("base_weight", "baseline", "baseline | probabilities: the baseline weight redistributed to the phone sample, or 1 / product of the probabilities (HFPS Vol. 2 pp. 10-11, 21)"),
    ("nonresponse", "propensity", "propensity | probit | classes | ebalance | ebalance_nobase | post | none (HFPS Vol. 2 Annex A2.3.1: propensity in deciles; ebalance and post are L2 program practice)"),
    ("covariates", "urban i.educ head_age assets hhsize", "baseline variables of the propensity model or of the balancing (i. for categories)"),
    ("classes", "region urban", "class variables for nonresponse classes (UN 2008 section 6.5)"),
    ("cells", "region urban head_female agegrp", "cell variables for nonresponse post: the columns of the Cells sheet before total"),
    ("ngroups", "10", "propensity groups (deciles)"),
    ("moments", "1", "balancing moments: 1 means, 2 adds variances, 3 adds skewness"),
    ("trim", "none", "none | winsor | cap (HFPS Vol. 2 p. 11 winsorises at the 98th percentile; UN 2008 section 6.8)"),
    ("trim_pct", "98", "winsorising percentile"),
    ("trim_cap", "", "absolute cap for trim cap, or blank"),
    ("trim_capmult", "", "cap as a multiple of the median, or blank"),
    ("calibrate", "rake", "none | rake | regress | post (Deville and Sarndal 1992; Deville, Sarndal and Sautory 1993)"),
    ("cal_vars", "region urban", "categorical variables to calibrate on; every category must be on the Totals sheet; blank = every variable there"),
    ("cal_ll", "", "lower bound on the weight ratio (regress), or blank"),
    ("cal_ul", "", "upper bound, or blank"),
    ("report_by", "region", "variable for the by-group table of the weighting report, or blank"),
    ("report_files", "yes", "yes | no: write the Word method note and the Excel report into output_dir"),
]

FILES_COLS = ["role", "path", "id", "response", "psu", "strata", "weight", "probabilities", "phone", "area_id", "round", "note"]
FILES_NOTES = {
    "role": "baseline | round | household_list | listing",
    "path": "file path, relative to this workbook's folder or absolute (.dta, .csv or .xlsx)",
    "id": "household identifier (baseline, round, household_list)",
    "response": "round: the 0/1 response indicator of that round",
    "psu": "baseline: the PSU variable (the phone subsample is drawn within it)",
    "strata": "baseline: the stratum variable (by-groups of the base weight and of the report)",
    "weight": "baseline: the baseline weight (base_weight baseline)",
    "probabilities": "baseline: the stage inclusion probabilities, space separated (base_weight probabilities)",
    "phone": "baseline: 1 = household in the phone sample (blank: every household)",
    "area_id": "household_list or listing: the area identifier that matches Frame area_id",
    "round": "round: the round number",
    "note": "free text",
}


def header(ws, cols, notes=None, required=()):
    for j, c in enumerate(cols, start=1):
        cell = ws.cell(row=1, column=j, value=c)
        cell.font = SUB
        cell.fill = REQ if c in required else FILL
        ws.column_dimensions[get_column_letter(j)].width = max(14, min(60, len(notes.get(c, "")) // 2 if notes else 14))
    if notes:
        for j, c in enumerate(cols, start=1):
            if c in notes:
                ws.cell(row=1, column=j).comment = None
    ws.freeze_panes = "A2"


def write_rows(ws, df, start=2):
    for i, row in enumerate(df.itertuples(index=False), start=start):
        for j, v in enumerate(row, start=1):
            if pd.isna(v):
                continue
            if hasattr(v, "item"):
                v = v.item()
            ws.cell(row=i, column=j, value=v)


def build(out, path_prefix):
    wb = Workbook()
    # ---- Readme ----
    w = wb.active
    w.title = "Readme"
    w.column_dimensions["A"].width = 120
    lines = [
        ("rtmtools input workbook (schema rtm-input-1.0)", HEAD),
        ("Fill in this workbook, then in Stata:  rtminput using rtm_inputs.xlsx, check   (validates every sheet and writes a Checks sheet into a copy)", None),
        ("                                      rtminput using rtm_inputs.xlsx, sampling   (size, frame check, allocation, draw, second stage, phone subsample)", None),
        ("                                      rtminput using rtm_inputs.xlsx, weights round(1)   (base, nonresponse, trim, calibrate, report for that round)", None),
        ("A folder of CSV files named design.csv, strata.csv, frame.csv, totals.csv, cells.csv, prefixes.csv, files.csv works the same (rtminput using folder/).", None),
        ("", None),
        ("Design: one row per setting (key, value). Blank = the default, whose source is in the note column. Rows whose key starts with # are headings.", SUB),
        ("Strata: one row per stratum (stratum_id, label, population_hh, population_persons). stratum_id must match Frame.", SUB),
        ("Frame: one row per area (area_id, stratum_id, size_hh required; size_persons, x, y, exclude and any sorting columns optional).", SUB),
        ("Totals: one row per category for calibration (variable, level, total), as the Stata command rtmweight calibrate reads it.", SUB),
        ("Cells: one row per census cell for cell post-stratification (the cell variables, then total).", SUB),
        ("Prefixes: one row per operator prefix for random digit dialing (prefix, n_numbers, digits).", SUB),
        ("Files: one row per household-level data file (baseline, round, household_list, listing) with the names of its key variables. The data stay in those files.", SUB),
        ("", None),
        ("The example rows are the package's synthetic data: nothing here is a real household. Replace every sheet with your program's inputs.", MUTED),
        ("Nothing in this workbook is a design recommendation: the choices are yours, and each default's source is named so you can check it.", MUTED),
        ("Column headings shaded yellow are required; the others may stay empty. Do not rename the sheets or the heading row.", MUTED),
    ]
    for i, (txt, font) in enumerate(lines, start=1):
        c = w.cell(row=i, column=1, value=txt)
        if font:
            c.font = font
    # ---- Design ----
    d = wb.create_sheet("Design")
    header(d, ["key", "value", "note"], required=("key", "value"))
    d.column_dimensions["A"].width = 22
    d.column_dimensions["B"].width = 36
    d.column_dimensions["C"].width = 120
    for i, (k, v, note) in enumerate(DESIGN, start=2):
        c = d.cell(row=i, column=1, value=k)
        if k.startswith("#"):
            c.font = SUB
            continue
        if v != "":
            d.cell(row=i, column=2, value=v)
        d.cell(row=i, column=3, value=note).font = MUTED
    # ---- Strata (from the synthetic frame) ----
    fr = pd.read_stata(os.path.join(DATA, "rtm_frame.dta"), convert_categoricals=False)
    frl = pd.read_stata(os.path.join(DATA, "rtm_frame.dta"))
    lab = frl.groupby("stratum").agg(region=("region", "first"), urban=("urban", "first")).reset_index()
    pop = fr.groupby("stratum").agg(population_hh=("N_str", "first")).reset_index()
    st = lab.merge(pop, on="stratum")
    st["label"] = st["region"].astype(str) + " " + st["urban"].astype(str)
    st = st.rename(columns={"stratum": "stratum_id"})[["stratum_id", "label", "population_hh"]]
    st["population_persons"] = None
    s = wb.create_sheet("Strata")
    header(s, ["stratum_id", "label", "population_hh", "population_persons"], required=("stratum_id",))
    write_rows(s, st)
    # ---- Frame ----
    f = wb.create_sheet("Frame")
    fdf = fr[["ea_id", "stratum", "ea_size_census", "region", "urban", "network"]].rename(
        columns={"ea_id": "area_id", "stratum": "stratum_id", "ea_size_census": "size_hh"})
    fdf["size_persons"] = None
    fdf["x"] = None
    fdf["y"] = None
    fdf["exclude"] = 0
    fdf = fdf[["area_id", "stratum_id", "size_hh", "size_persons", "region", "urban", "network", "x", "y", "exclude"]]
    header(f, list(fdf.columns), required=("area_id", "stratum_id", "size_hh"))
    write_rows(f, fdf)
    # ---- Totals ----
    t = wb.create_sheet("Totals")
    tdf = pd.read_stata(os.path.join(DATA, "rtm_totals_long.dta"), convert_categoricals=False)
    header(t, ["variable", "level", "total"], required=("variable", "level", "total"))
    write_rows(t, tdf[["variable", "level", "total"]])
    # ---- Cells ----
    c = wb.create_sheet("Cells")
    cdf = pd.read_stata(os.path.join(DATA, "rtm_totals_cells.dta"), convert_categoricals=False).rename(columns={"N_hh": "total"})
    header(c, list(cdf.columns), required=("total",))
    write_rows(c, cdf)
    # ---- Prefixes ----
    p = wb.create_sheet("Prefixes")
    header(p, ["prefix", "n_numbers", "digits"], required=("prefix", "n_numbers", "digits"))
    for i, (pre, n, dg) in enumerate([(61, 300, 6), (62, 300, 6), (77, 400, 6)], start=2):
        p.cell(row=i, column=1, value=pre)
        p.cell(row=i, column=2, value=n)
        p.cell(row=i, column=3, value=dg)
    # ---- Files ----
    fl = wb.create_sheet("Files")
    header(fl, FILES_COLS, required=("role", "path"))
    fl.column_dimensions["B"].width = 34
    fl.column_dimensions["L"].width = 60
    rows = [
        {"role": "baseline", "path": path_prefix + "rtm_baseline.dta", "id": "hhid", "psu": "psu", "strata": "stratum", "weight": "w_base",
         "probabilities": "pi_ea pi_hh", "phone": "phone", "note": "the baseline sample, one row per household"},
        {"role": "round", "path": path_prefix + "rtm_round1.dta", "id": "hhid", "response": "responded_r1", "round": 1,
         "note": "round 1: one row per phone-sample household with the response indicator"},
        {"role": "household_list", "path": path_prefix + "rtm_population.dta", "id": "hhid_pop", "area_id": "ea_id",
         "note": "household list of the areas for the second stage (one row per household)"},
    ]
    for i, r in enumerate(rows, start=2):
        for j, col in enumerate(FILES_COLS, start=1):
            if col in r:
                fl.cell(row=i, column=j, value=r[col])
    # the Files columns, explained on the Readme sheet (the Files sheet holds data rows only)
    r0 = len(lines) + 2
    w.cell(row=r0, column=1, value="Files sheet columns").font = SUB
    for k, (col, note) in enumerate(FILES_NOTES.items(), start=r0 + 1):
        w.cell(row=k, column=1, value=f"{col}: {note}").font = MUTED
    wb.save(out)
    return out


if __name__ == "__main__":
    print("written:", build(os.path.join(ROOT, "rtm_inputs.xlsx"), ""))
    print("written:", build(os.path.join(ROOT, "portable_do_files", "rtm_inputs.xlsx"), "template/"))
